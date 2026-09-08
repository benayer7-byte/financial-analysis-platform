import pandas as pd
from openpyxl import load_workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.chart.layout import Layout, ManualLayout

company_name = "Apple"
ticker = "AAPL"
csv_path = f"data/raw/{ticker.lower()}.csv"
excel_path = f"data/processed/{ticker.lower()}_info.xlsx"

df = pd.read_csv(csv_path)
wb = load_workbook(excel_path)

if f"{company_name} Charts" in wb.sheetnames:
    del wb[f"{company_name} Charts"]
ws = wb.create_sheet(f"{company_name} Charts")
ws["A1"] = "Fiscal Year"
ws["B1"] = "Revenue"

for row, (_, data) in enumerate(df.iterrows(), start=2):
    ws[f"A{row}"] = data["fiscal_year"]
    ws[f"B{row}"] = data["revenue"] / 1000

revenue_chart = LineChart()
revenue_chart.title = f"{company_name} Revenue Over Time"
revenue_chart.style = 2
revenue_chart.varyColors = False

date = Reference(ws, min_col=1, min_row=2, max_row=len(df) + 1)
revenue = Reference(ws, min_col=2, min_row=1, max_row=len(df) + 1)
revenue_chart.add_data(revenue, titles_from_data=True)
revenue_chart.set_categories(date)

series = revenue_chart.series[0]
series.marker.symbol = "circle"
series.marker.size = 7
series.smooth = False

line_color = "1F77B4"
series.graphicalProperties.line.solidFill = line_color
series.graphicalProperties.line.width = 22000
series.marker.graphicalProperties.solidFill = line_color
series.marker.graphicalProperties.line.solidFill = line_color

revenue_chart.x_axis.title = "Fiscal Year"
revenue_chart.y_axis.title = "Revenue ($ Billions)"
revenue_chart.x_axis.axPos = "b"   # bottom axis (openpyxl defaults this to "l", which is wrong)
revenue_chart.y_axis.axPos = "l"   # left axis
revenue_chart.height = 10
revenue_chart.width = 20
revenue_chart.legend = None

vals = df["revenue"] / 1000
pad = (vals.max() - vals.min()) * 0.15 or 5
revenue_chart.y_axis.scaling.min = round(vals.min() - pad, -1)
revenue_chart.y_axis.scaling.max = round(vals.max() + pad, -1)
revenue_chart.y_axis.majorUnit = 10

# --- make axis lines + value labels visible ---
for axis in (revenue_chart.x_axis, revenue_chart.y_axis):
    axis.delete = False                 # make sure axis isn't hidden
    axis.majorTickMark = "out"          # tick marks on the axis line
    axis.minorTickMark = "none"
    axis.tickLblPos = "nextTo"          # show the value labels next to axis
    axis.spPr = None  # reset first
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
axis_line_props = GraphicalProperties()
axis_line_props.line = LineProperties(solidFill="808080", w=9525)  # thin grey axis line
revenue_chart.x_axis.spPr = axis_line_props
# reuse a fresh instance for y-axis (can't share the same object)
axis_line_props_y = GraphicalProperties()
axis_line_props_y.line = LineProperties(solidFill="808080", w=9525)
revenue_chart.y_axis.spPr = axis_line_props_y

# lighter gridlines so they don't visually compete with axis/title
from openpyxl.chart.axis import ChartLines
gridline_props = GraphicalProperties()
gridline_props.line = LineProperties(solidFill="D9D9D9", w=6350)
revenue_chart.y_axis.majorGridlines = ChartLines(spPr=gridline_props)

# reserve room so gridlines/plot area never run into the title
revenue_chart.layout = Layout(
    manualLayout=ManualLayout(
        x=0.12, y=0.20, w=0.82, h=0.68,
        xMode="edge", yMode="edge"
    )
)

ws.add_chart(revenue_chart, "G1")

#first chart complete, now for the second chart 
#-------------
def build_metric_chart(title, column_name, y_axis_title, data_col, color, divide_by_1000=True):
    col_letter = chr(ord('A') + data_col - 1)
    ws[f"{col_letter}1"] = column_name

    for row, (_, data) in enumerate(df.iterrows(), start=2):
        value = data[column_name]
        ws[f"{col_letter}{row}"] = value / 1000 if divide_by_1000 else value

    chart = LineChart()
    chart.title = title
    chart.style = 2

    cats = Reference(ws, min_col=1, min_row=2, max_row=len(df) + 1)
    data_ref = Reference(ws, min_col=data_col, min_row=1, max_row=len(df) + 1)
    chart.add_data(data_ref, titles_from_data=True)
    chart.set_categories(cats)

    series = chart.series[0]
    series.marker.symbol = "circle"
    series.marker.size = 7
    series.smooth = False
    series.graphicalProperties.line.solidFill = color
    series.graphicalProperties.line.width = 22000
    series.marker.graphicalProperties.solidFill = color
    series.marker.graphicalProperties.line.solidFill = color

    chart.x_axis.title = "Fiscal Year"
    chart.y_axis.title = y_axis_title
    for axis in (chart.x_axis, chart.y_axis):
        axis.delete = False
        axis.tickLblPos = "nextTo"
    chart.height = 10
    chart.width = 20
    chart.legend = None
    chart.layout = Layout(manualLayout=ManualLayout(x=0.12, y=0.20, w=0.82, h=0.68, xMode="edge", yMode="edge"))
    return chart

net_income_chart = build_metric_chart(f"{company_name} Net Income Over Time", "net_income", "Net Income ($ Billions)", 3, "2CA02C")
eps_chart = build_metric_chart(f"{company_name} EPS Over Time", "diluted_eps", "EPS ($)", 4, "9467BD", divide_by_1000=False)
stock_price_chart = build_metric_chart(f"{company_name} Stock Price Over Time", "stock_price", "Stock Price ($)", 5, "FF7F0E", divide_by_1000=False)

ws.add_chart(net_income_chart, "G20")
ws.add_chart(eps_chart, "G39")
ws.add_chart(stock_price_chart, "G58")

#Second Chart done, now add third chart 








wb.save(excel_path)
print("saved")

