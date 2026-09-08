import pandas as pd
import numpy as np
from openpyxl import load_workbook

company_name = "Apple"
ticker = "AAPL"
csv_path = f"data/raw/{ticker.lower()}.csv"
excel_path = f"data/processed/{ticker.lower()}_info.xlsx"

df = pd.read_csv(csv_path)

# --- Method 1: Average growth rate ---
growth_rates = df["revenue"].pct_change().dropna()
avg_growth_rate = growth_rates.mean()

last_revenue = df["revenue"].iloc[-1]
avg_growth_forecast = [
    last_revenue * (1 + avg_growth_rate) ** 1,
    last_revenue * (1 + avg_growth_rate) ** 2,
]

# --- Method 2: Linear trend ---
years = df["fiscal_year"].values
revenue = df["revenue"].values
slope, intercept = np.polyfit(years, revenue, 1)

future_years = [years[-1] + 1, years[-1] + 2]
linear_forecast = [slope * year + intercept for year in future_years]

print("Average growth rate:", round(avg_growth_rate * 100, 2), "%")
print("Average-growth forecast (2026, 2027):", [round(v) for v in avg_growth_forecast])
print("Linear trend forecast (2026, 2027):", [round(v) for v in linear_forecast])

sheet_name = f"{company_name} Forecasting"

from openpyxl.chart import LineChart, Reference
from openpyxl.chart.layout import Layout, ManualLayout

# Build one combined table: actual years + both forecast methods
all_years = list(df["fiscal_year"]) + future_years
avg_growth_series = [v / 1000 for v in list(df["revenue"]) + avg_growth_forecast]
linear_series = [v / 1000 for v in list(df["revenue"]) + linear_forecast]

sheet_name = f"{company_name} Forecasting"

wb = load_workbook(excel_path)
if sheet_name in wb.sheetnames:
    del wb[sheet_name]
ws = wb.create_sheet(sheet_name)

ws["A1"] = "Fiscal Year"
ws["B1"] = "Avg Growth Forecast"
ws["C1"] = "Linear Trend Forecast"

for row, year in enumerate(all_years, start=2):
    ws[f"A{row}"] = year
    ws[f"B{row}"] = round(avg_growth_series[row - 2], 1)
    ws[f"C{row}"] = round(linear_series[row - 2], 1)

def build_chart(title, data_col, color):
    chart = LineChart()
    chart.title = title
    chart.style = 2

    cats = Reference(ws, min_col=1, min_row=2, max_row=len(all_years) + 1)
    data = Reference(ws, min_col=data_col, min_row=1, max_row=len(all_years) + 1)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)

    series = chart.series[0]
    series.marker.symbol = "circle"
    series.marker.size = 7
    series.smooth = False
    series.graphicalProperties.line.solidFill = color
    series.graphicalProperties.line.width = 22000
    series.marker.graphicalProperties.solidFill = color
    series.marker.graphicalProperties.line.solidFill = color

    chart.x_axis.title = "Fiscal Year"
    chart.y_axis.title = "Revenue ($ Billions)"
    chart.height = 10
    chart.width = 20
    chart.legend = None
    chart.layout = Layout(manualLayout=ManualLayout(x=0.12, y=0.20, w=0.82, h=0.68, xMode="edge", yMode="edge"))
    for axis in (chart.x_axis, chart.y_axis):
        axis.delete = False
        axis.tickLblPos = "nextTo"
    return chart

avg_chart = build_chart(f"{company_name} Revenue Forecast (Avg Growth Rate)", 2, "1F77B4")
linear_chart = build_chart(f"{company_name} Revenue Forecast (Linear Trend)", 3, "D62728")

ws.add_chart(avg_chart, "E1")
ws.add_chart(linear_chart, "E20")

wb.save(excel_path)
print("saved")
