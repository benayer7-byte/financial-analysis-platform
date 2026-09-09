import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.table import Table, TableStyleInfo

company_name = "Apple"
ticker = "AAPL"
import os
script_dir = os.path.dirname(os.path.abspath(__file__))
project_root = os.path.dirname(script_dir)

csv_path = os.path.join(project_root, "data", "raw", f"{ticker.lower()}.csv")
excel_path = os.path.join(project_root, "data", "processed", f"{ticker.lower()}_info.xlsx")

import sqlite3

db_path = os.path.join(project_root, "data", "financials.db")
conn = sqlite3.connect(db_path)
df = pd.read_sql(f"SELECT * FROM financials WHERE ticker = '{ticker}'", conn)
conn.close()

# ============================================================
# RATIOS
# ============================================================

df["current_Ratio"] = df['current_assets'] / df['current_liabilities']
df["net_profit_margin"] = df['net_income'] / df['revenue']
df["roa"] = df['net_income'] / df['total_assets']
df["roe"] = df['net_income'] / df['shareholders_equity']
df["debt_to_equity_ratio"] = df['total_liabilities'] / df['shareholders_equity']
df["operating_cash_flow_ratio"] = df['operating_cash_flow'] / df['current_liabilities']

ratio_table = df[["fiscal_year", "current_Ratio", "net_profit_margin", "roa", "roe",
                  "debt_to_equity_ratio", "operating_cash_flow_ratio"]]
ratio_table = ratio_table.set_index("fiscal_year").T
ratio_table.index = ["Current Ratio", "Net Profit Margin", "Return on Assets (ROA)",
                     "Return on Equity (ROE)", "Debt to Equity Ratio", "Operating Cash Flow Ratio"]

for ratio in ["Net Profit Margin", "Return on Assets (ROA)", "Return on Equity (ROE)"]:
    ratio_table.loc[ratio] = ratio_table.loc[ratio] * 100
ratio_table = ratio_table.round(2)

ratio_table.to_excel(excel_path, sheet_name=f"{company_name} Ratios")

wb = load_workbook(excel_path)
ws = wb.active
ws.title = f"{company_name} Ratios"

for column in ws.columns:
    max_length = 0
    column_letter = column[0].column_letter
    for cell in column:
        if cell.value is not None:
            max_length = max(max_length, len(str(cell.value)))
    ws.column_dimensions[column_letter].width = max_length + 2

ws["A1"] = f"Ratios for {company_name} Inc."
ws["A1"].font = Font(size=14, bold=True, color="1F4E78")
ws["A1"].alignment = Alignment(horizontal="center")
ws["A1"].fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")

for cell in ws[1][1:]:
    cell.font = Font(size=12, bold=True)
for cell in ws["A"][1:]:
    cell.font = Font(size=12, bold=True)

thin_border = Border(left=Side(style='thin'), right=Side(style='thin'),
                     top=Side(style='thin'), bottom=Side(style='thin'))

for row in ws.iter_rows():
    for cell in row:
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center", vertical="center")

for cell in ws[1]:
    cell.value = str(cell.value)

table = Table(displayName=f"{ticker}Ratios", ref=f"A1:E{ws.max_row}")
style = TableStyleInfo(name="TableStyleMedium9", showFirstColumn=False,
                       showLastColumn=False, showRowStripes=True, showColumnStripes=False)
table.tableStyleInfo = style

for sheet in wb.worksheets:
    if f"{ticker}Ratios" in sheet.tables:
        del sheet.tables[f"{ticker}Ratios"]
ws.add_table(table)

percentage_ratios = ["Net Profit Margin", "Return on Assets (ROA)", "Return on Equity (ROE)"]

for row in ws.iter_rows(min_row=2):
    ratio_name = row[0].value
    for cell in row[1:]:
        if ratio_name in percentage_ratios:
            cell.number_format = '0.00"%"'
        else:
            cell.number_format = "0.00"

wb.save(excel_path)
print("Ratios done")

# ============================================================
# CHARTS
# ============================================================

from openpyxl.chart import LineChart, Reference
from openpyxl.chart.layout import Layout, ManualLayout

wb = load_workbook(excel_path)

sheet_name = f"{company_name} Charts"
if sheet_name in wb.sheetnames:
    del wb[sheet_name]
ws = wb.create_sheet(sheet_name)

ws["A1"] = "Fiscal Year"
ws["B1"] = "Revenue"

for row, (_, data) in enumerate(df.iterrows(), start=2):
    ws[f"A{row}"] = data["fiscal_year"]
    ws[f"B{row}"] = data["revenue"] / 1000

revenue_chart = LineChart()
revenue_chart.title = f"{company_name} Revenue Over Time"
revenue_chart.style = 2
revenue_chart.varyColors = False

cats = Reference(ws, min_col=1, min_row=2, max_row=len(df) + 1)
data_ref = Reference(ws, min_col=2, min_row=1, max_row=len(df) + 1)
revenue_chart.add_data(data_ref, titles_from_data=True)
revenue_chart.set_categories(cats)

series = revenue_chart.series[0]
series.marker.symbol = "circle"
series.marker.size = 7
series.smooth = False

line_color = "1F77B4"
series.graphicalProperties.line.solidFill = line_color
series.graphicalProperties.line.width = 22000
series.marker.graphicalProperties.solidFill = line_color
series.marker.graphicalProperties.line.solidFill = line_color

revenue_chart.x_axis.title = "Fiscal Year"
revenue_chart.y_axis.title = "Revenue ($ Billions)"
revenue_chart.x_axis.axPos = "b"
revenue_chart.y_axis.axPos = "l"
revenue_chart.height = 10
revenue_chart.width = 20
revenue_chart.legend = None

vals = df["revenue"] / 1000
pad = (vals.max() - vals.min()) * 0.15 or 5
revenue_chart.y_axis.scaling.min = round(vals.min() - pad, -1)
revenue_chart.y_axis.scaling.max = round(vals.max() + pad, -1)
revenue_chart.y_axis.majorUnit = 10

for axis in (revenue_chart.x_axis, revenue_chart.y_axis):
    axis.delete = False
    axis.majorTickMark = "out"
    axis.minorTickMark = "none"
    axis.tickLblPos = "nextTo"

from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
axis_line_props = GraphicalProperties()
axis_line_props.line = LineProperties(solidFill="808080", w=9525)
revenue_chart.x_axis.spPr = axis_line_props
axis_line_props_y = GraphicalProperties()
axis_line_props_y.line = LineProperties(solidFill="808080", w=9525)
revenue_chart.y_axis.spPr = axis_line_props_y

from openpyxl.chart.axis import ChartLines
gridline_props = GraphicalProperties()
gridline_props.line = LineProperties(solidFill="D9D9D9", w=6350)
revenue_chart.y_axis.majorGridlines = ChartLines(spPr=gridline_props)

revenue_chart.layout = Layout(
    manualLayout=ManualLayout(x=0.12, y=0.20, w=0.82, h=0.68, xMode="edge", yMode="edge")
)

ws.add_chart(revenue_chart, "G1")

def build_metric_chart(title, column_name, y_axis_title, data_col, color, divide_by_1000=True):
    col_letter = chr(ord('A') + data_col - 1)
    ws[f"{col_letter}1"] = column_name

    for row, (_, data) in enumerate(df.iterrows(), start=2):
        value = data[column_name]
        ws[f"{col_letter}{row}"] = value / 1000 if divide_by_1000 else value

    chart = LineChart()
    chart.title = title
    chart.style = 2

    cats = Reference(ws, min_col=1, min_row=2, max_row=len(df) + 1)
    data_ref = Reference(ws, min_col=data_col, min_row=1, max_row=len(df) + 1)
    chart.add_data(data_ref, titles_from_data=True)
    chart.set_categories(cats)

    series = chart.series[0]
    series.marker.symbol = "circle"
    series.marker.size = 7
    series.smooth = False
    series.graphicalProperties.line.solidFill = color
    series.graphicalProperties.line.width = 22000
    series.marker.graphicalProperties.solidFill = color
    series.marker.graphicalProperties.line.solidFill = color

    chart.x_axis.title = "Fiscal Year"
    chart.y_axis.title = y_axis_title
    for axis in (chart.x_axis, chart.y_axis):
        axis.delete = False
        axis.tickLblPos = "nextTo"
    chart.height = 10
    chart.width = 20
    chart.legend = None
    chart.layout = Layout(manualLayout=ManualLayout(x=0.12, y=0.20, w=0.82, h=0.68, xMode="edge", yMode="edge"))
    return chart

net_income_chart = build_metric_chart(f"{company_name} Net Income Over Time", "net_income", "Net Income ($ Billions)", 3, "2CA02C")
eps_chart = build_metric_chart(f"{company_name} EPS Over Time", "diluted_eps", "EPS ($)", 4, "9467BD", divide_by_1000=False)
stock_price_chart = build_metric_chart(f"{company_name} Stock Price Over Time", "stock_price", "Stock Price ($)", 5, "FF7F0E", divide_by_1000=False)

ws.add_chart(net_income_chart, "G20")
ws.add_chart(eps_chart, "G39")
ws.add_chart(stock_price_chart, "G58")

wb.save(excel_path)
print("Charts done")

# ============================================================
# VALUATION
# ============================================================

market_cap = df["stock_price"] * df["diluted_weighted_average_shares"]
free_cash_flow = df["operating_cash_flow"] - df["capital_expenditures"]

df["pe_ratio"] = df["stock_price"] / df["diluted_eps"]
df["ps_ratio"] = market_cap / df["revenue"]
df["price_to_book"] = market_cap / df["shareholders_equity"]
df["fcf_yield"] = free_cash_flow / market_cap

valuation_table = df[["fiscal_year", "pe_ratio", "ps_ratio", "price_to_book", "fcf_yield"]]
valuation_table = valuation_table.set_index("fiscal_year").T

valuation_table.index = ["P/E Ratio", "P/S Ratio", "Price-to-Book", "FCF Yield"]
valuation_table = valuation_table.round(2)
valuation_table.loc["FCF Yield"] = valuation_table.loc["FCF Yield"] * 100

sheet_name = f"{company_name} Valuation"

with pd.ExcelWriter(excel_path, engine="openpyxl", mode="a", if_sheet_exists="replace") as writer:
    valuation_table.to_excel(writer, sheet_name=sheet_name)

wb = load_workbook(excel_path)
ws = wb[sheet_name]

for column in ws.columns:
    max_length = 0
    column_letter = column[0].column_letter
    for cell in column[1:]:
        if cell.value is not None:
            max_length = max(max_length, len(str(cell.value)))
    ws.column_dimensions[column_letter].width = max_length + 2

ws["A1"] = f"Valuation for {company_name} Inc."
ws.column_dimensions["A"].width = max(ws.column_dimensions["A"].width, len(ws["A1"].value) + 2)
ws["A1"].font = Font(size=14, bold=True, color="1F4E78")
ws["A1"].alignment = Alignment(horizontal="center")
ws["A1"].fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")

for cell in ws[1][1:]:
    cell.font = Font(size=12, bold=True)
for cell in ws["A"][1:]:
    cell.font = Font(size=12, bold=True)

for row in ws.iter_rows():
    for cell in row:
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center", vertical="center")

for cell in ws[1]:
    cell.value = str(cell.value)

table = Table(displayName=f"{ticker}Valuation", ref=f"A1:E{ws.max_row}")
style = TableStyleInfo(name="TableStyleMedium9", showFirstColumn=False,
                       showLastColumn=False, showRowStripes=True, showColumnStripes=False)
table.tableStyleInfo = style

for sheet in wb.worksheets:
    if f"{ticker}Valuation" in sheet.tables:
        del sheet.tables[f"{ticker}Valuation"]
ws.add_table(table)

for row in ws.iter_rows(min_row=2):
    ratio_name = row[0].value
    for cell in row[1:]:
        if ratio_name == "FCF Yield":
            cell.number_format = '0.00"%"'
        else:
            cell.number_format = "0.00"

wb.save(excel_path)
print("Valuation done")

# ============================================================
# FORECASTING
# ============================================================

import numpy as np

growth_rates = df["revenue"].pct_change().dropna()
avg_growth_rate = growth_rates.mean()

last_revenue = df["revenue"].iloc[-1]
avg_growth_forecast = [
    last_revenue * (1 + avg_growth_rate) ** 1,
    last_revenue * (1 + avg_growth_rate) ** 2,
]

years = df["fiscal_year"].values
revenue = df["revenue"].values
slope, intercept = np.polyfit(years, revenue, 1)

future_years = [years[-1] + 1, years[-1] + 2]
linear_forecast = [slope * year + intercept for year in future_years]

all_years = list(df["fiscal_year"]) + future_years
avg_growth_series = [v / 1000 for v in list(df["revenue"]) + avg_growth_forecast]
linear_series = [v / 1000 for v in list(df["revenue"]) + linear_forecast]

sheet_name = f"{company_name} Forecasting"

wb = load_workbook(excel_path)
if sheet_name in wb.sheetnames:
    del wb[sheet_name]
ws = wb.create_sheet(sheet_name)

ws["A1"] = "Fiscal Year"
ws["B1"] = "Avg Growth Forecast"
ws["C1"] = "Linear Trend Forecast"

for row, year in enumerate(all_years, start=2):
    ws[f"A{row}"] = year
    ws[f"B{row}"] = round(avg_growth_series[row - 2], 1)
    ws[f"C{row}"] = round(linear_series[row - 2], 1)

def build_chart(title, data_col, color):
    chart = LineChart()
    chart.title = title
    chart.style = 2

    cats = Reference(ws, min_col=1, min_row=2, max_row=len(all_years) + 1)
    data_ref = Reference(ws, min_col=data_col, min_row=1, max_row=len(all_years) + 1)
    chart.add_data(data_ref, titles_from_data=True)
    chart.set_categories(cats)

    series = chart.series[0]
    series.marker.symbol = "circle"
    series.marker.size = 7
    series.smooth = False
    series.graphicalProperties.line.solidFill = color
    series.graphicalProperties.line.width = 22000
    series.marker.graphicalProperties.solidFill = color
    series.marker.graphicalProperties.line.solidFill = color

    chart.x_axis.title = "Fiscal Year"
    chart.y_axis.title = "Revenue ($ Billions)"
    for axis in (chart.x_axis, chart.y_axis):
        axis.delete = False
        axis.tickLblPos = "nextTo"
    chart.height = 10
    chart.width = 20
    chart.legend = None
    chart.layout = Layout(manualLayout=ManualLayout(x=0.12, y=0.20, w=0.82, h=0.68, xMode="edge", yMode="edge"))
    return chart

avg_chart = build_chart(f"{company_name} Revenue Forecast (Avg Growth Rate)", 2, "1F77B4")
linear_chart = build_chart(f"{company_name} Revenue Forecast (Linear Trend)", 3, "D62728")

ws.add_chart(avg_chart, "E1")
ws.add_chart(linear_chart, "E20")

wb.save(excel_path)
print("Forecasting done")

# ============================================================
# SUMMARY
# ============================================================

market_cap = df["stock_price"] * df["diluted_weighted_average_shares"]
free_cash_flow_summary = df["operating_cash_flow"] - df["capital_expenditures"]

df["revenue_growth"] = df["revenue"].pct_change() * 100
df["net_margin"] = (df["net_income"] / df["revenue"]) * 100
df["roe"] = (df["net_income"] / df["shareholders_equity"]) * 100
df["fcf"] = free_cash_flow_summary / 1000
df["debt_to_equity"] = df["total_liabilities"] / df["shareholders_equity"]
df["pe_ratio"] = df["stock_price"] / df["diluted_eps"]
df["current_ratio"] = df["current_assets"] / df["current_liabilities"]

kpi_table = df[["fiscal_year", "revenue_growth", "net_margin", "roe", "fcf", "debt_to_equity", "pe_ratio"]]
kpi_table = kpi_table.round(2)
kpi_table = kpi_table.set_index("fiscal_year").T
kpi_table.index = ["Revenue Growth", "Net Margin", "ROE", "Free Cash Flow ($B)", "Debt to Equity", "P/E Ratio"]

current_values = kpi_table[2025]
trend_direction = []
for metric in kpi_table.index:
    first_value = kpi_table.loc[metric].iloc[0]
    last_value = kpi_table.loc[metric].iloc[-1]
    if last_value > first_value:
        trend_direction.append("↑")
    elif last_value < first_value:
        trend_direction.append("↓")
    else:
        trend_direction.append("→")
kpi_table["Trend"] = trend_direction

sheet_name = f"{company_name} Summary"

with pd.ExcelWriter(excel_path, engine="openpyxl", mode="a", if_sheet_exists="replace") as writer:
    kpi_table.to_excel(writer, sheet_name=sheet_name)

wb = load_workbook(excel_path)
ws = wb[sheet_name]

for column in ws.columns:
    max_length = 0
    column_letter = column[0].column_letter
    for cell in column[1:]:
        if cell.value is not None:
            max_length = max(max_length, len(str(cell.value)))

ws["A1"] = f"{company_name} Financial Summary"
ws.column_dimensions["A"].width = len(ws["A1"].value) + 6
ws["A1"].font = Font(size=14, bold=True, color="1F4E78")
ws["A1"].alignment = Alignment(horizontal="center")
ws["A1"].fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")

for cell in ws[1][1:]:
    cell.font = Font(size=12, bold=True)
for cell in ws["A"][1:]:
    cell.font = Font(size=12, bold=True)

for row in ws.iter_rows():
    for cell in row:
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center", vertical="center")

for cell in ws[1]:
    cell.value = str(cell.value)

table = Table(displayName=f"{ticker}Summary", ref=f"A1:F{ws.max_row}")
style = TableStyleInfo(name="TableStyleMedium9", showFirstColumn=False,
                       showLastColumn=False, showRowStripes=True, showColumnStripes=False)
table.tableStyleInfo = style
for sheet in wb.worksheets:
    if f"{ticker}Summary" in sheet.tables:
        del sheet.tables[f"{ticker}Summary"]
ws.add_table(table)

percentage_rows = ["Revenue Growth", "Net Margin", "ROE"]
for row in ws.iter_rows(min_row=2):
    metric_name = row[0].value
    for cell in row[1:-1]:
        if metric_name in percentage_rows:
            cell.number_format = '0.00"%"'
        else:
            cell.number_format = "0.00"

def score_percent(value, bands=(25, 20, 15, 10)):
    if value > bands[0]: return 10, f"> {bands[0]}%"
    if value > bands[1]: return 8, f"{bands[1]}-{bands[0]}%"
    if value > bands[2]: return 6, f"{bands[2]}-{bands[1]}%"
    if value > bands[3]: return 4, f"{bands[3]}-{bands[2]}%"
    return 2, f"< {bands[3]}%"

def score_ratio_low_is_good(value, bands=(0.5, 1, 2, 3)):
    if value < bands[0]: return 10, f"< {bands[0]}"
    if value < bands[1]: return 8, f"{bands[0]}-{bands[1]}"
    if value < bands[2]: return 6, f"{bands[1]}-{bands[2]}"
    if value < bands[3]: return 4, f"{bands[2]}-{bands[3]}"
    return 2, f"> {bands[3]}"

def score_ratio_high_is_good(value, bands=(2, 1.5, 1, 0.5)):
    if value > bands[0]: return 10, f"> {bands[0]}"
    if value > bands[1]: return 8, f"{bands[1]}-{bands[0]}"
    if value > bands[2]: return 6, f"{bands[2]}-{bands[1]}"
    if value > bands[3]: return 4, f"{bands[3]}-{bands[2]}"
    return 2, f"< {bands[3]}"

fcf_margin_2025 = (df["fcf"].iloc[-1] / (df["revenue"].iloc[-1] / 1000)) * 100

metrics_2025 = {
    "Profitability (Net Margin)": (df["net_margin"].iloc[-1], score_percent),
    "Profitability (ROE)": (df["roe"].iloc[-1], score_percent),
    "Cash Flow (FCF Margin)": (fcf_margin_2025, score_percent),
    "Growth (Revenue Growth)": (df["revenue_growth"].iloc[-1], lambda v: score_percent(v, bands=(10, 5, 0, -5))),
    "Financial Risk (Debt to Equity)": (df["debt_to_equity"].iloc[-1], score_ratio_low_is_good),
    "Valuation (P/E Ratio)": (df["pe_ratio"].iloc[-1], lambda v: score_ratio_low_is_good(v, bands=(15, 20, 30, 40))),
    "Liquidity (Current Ratio)": (df["current_ratio"].iloc[-1], score_ratio_high_is_good),
}

profitability_score = (metrics_2025["Profitability (Net Margin)"][1](metrics_2025["Profitability (Net Margin)"][0])[0] +
                        metrics_2025["Profitability (ROE)"][1](metrics_2025["Profitability (ROE)"][0])[0]) / 2
cash_flow_score = metrics_2025["Cash Flow (FCF Margin)"][1](metrics_2025["Cash Flow (FCF Margin)"][0])[0]
growth_score = metrics_2025["Growth (Revenue Growth)"][1](metrics_2025["Growth (Revenue Growth)"][0])[0]
risk_score = metrics_2025["Financial Risk (Debt to Equity)"][1](metrics_2025["Financial Risk (Debt to Equity)"][0])[0]
valuation_score = metrics_2025["Valuation (P/E Ratio)"][1](metrics_2025["Valuation (P/E Ratio)"][0])[0]
liquidity_score = metrics_2025["Liquidity (Current Ratio)"][1](metrics_2025["Liquidity (Current Ratio)"][0])[0]

overall_score = (
    profitability_score * 0.20 + cash_flow_score * 0.20 + growth_score * 0.20 +
    risk_score * 0.15 + valuation_score * 0.15 + liquidity_score * 0.10
)

start_row = ws.max_row + 2
headers = ["Category", "2025 Value", "Score"]
for col, header in enumerate(headers, start=1):
    cell = ws.cell(row=start_row, column=col, value=header)
    cell.font = Font(size=12, bold=True)
    cell.fill = PatternFill(start_color="9DC3E6", end_color="9DC3E6", fill_type="solid")
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = thin_border

row_num = start_row + 1
for label, (value, scorer) in metrics_2025.items():
    score, _ = scorer(value)
    ws.cell(row=row_num, column=1, value=label)
    ws.cell(row=row_num, column=2, value=round(value, 2))
    ws.cell(row=row_num, column=3, value=score)
    for col in range(1, 4):
        cell = ws.cell(row=row_num, column=col)
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center", vertical="center")
    row_num += 1

ws.cell(row=row_num, column=1, value="Overall Score")
ws.cell(row=row_num, column=3, value=f"{round(overall_score, 1)}/10")
for col in range(1, 4):
    cell = ws.cell(row=row_num, column=col)
    cell.border = thin_border
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.font = Font(size=12, bold=True, color="1F4E78")

ws.column_dimensions["A"].width = 35
ws.column_dimensions["B"].width = 15

wb.save(excel_path)
print("Summary done")

print("\nAll sections complete! Full workbook built successfully.")

# ============================================================
# REORDER SHEETS
# ============================================================

wb = load_workbook(excel_path)

desired_order = [
    f"{company_name} Summary",
    f"{company_name} Charts",
    f"{company_name} Forecasting",
    f"{company_name} Valuation",
    f"{company_name} Ratios",
]

wb._sheets = [wb[name] for name in desired_order]
wb.save(excel_path)
print("Sheets reordered")